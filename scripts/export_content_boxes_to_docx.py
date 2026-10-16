#!/usr/bin/env python3
"""Export main informational HTML content boxes into a Word document.

The script scans html/*.html, extracts textual content from the main content area
("main.content" and its direct article/section descendants), and writes a review
DOCX for collaborative text verification.
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Iterable

from bs4 import BeautifulSoup
from docx import Document
from docx.shared import Pt


PROJECT_ROOT = Path(__file__).resolve().parents[1]
HTML_DIR = PROJECT_ROOT / "html"
OUTPUT_DOCX = PROJECT_ROOT / "docs" / "content" / "content-review-boxes.docx"
OUTPUT_DOCX_BY_TOPIC = PROJECT_ROOT / "docs" / "content" / "content-review-boxes-by-topic-generated.docx"
WEBSITE_URL = "http://tsvrudow.de/foerderverein/"

TOPIC_MAP: dict[str, str] = {
    "index": "Verein und Startseite",
    "about": "Verein und Startseite",
    "transparenz": "Transparenz und Organisation",
    "ansprechpartner": "Ansprechpartner und Kontakt",
    "aktuelles": "Aktuelles und Mitteilungen",
    "ergebnisse": "Sportbetrieb und Ergebnisse",
    "partner": "Partner und Sponsoring",
    "satzung": "Satzung und Rechtsgrundlagen",
    "impressum": "Rechtliches und Impressum",
}


def clean_line(text: str) -> str:
    return " ".join(text.split()).strip()


def is_inside_nav(element) -> bool:
    return element.find_parent("nav") is not None


def extract_box_blocks(box, box_title: str) -> list[tuple[str, str]]:
    blocks: list[tuple[str, str]] = []

    for element in box.find_all(["h1", "h2", "h3", "h4", "p", "ul", "ol", "table"]):
        if is_inside_nav(element):
            continue

        if element.name in {"h1", "h2", "h3", "h4"}:
            value = clean_line(element.get_text(" ", strip=True))
            if value and value != box_title:
                blocks.append(("subheading", value))
            continue

        if element.name == "p":
            value = clean_line(element.get_text(" ", strip=True))
            if value:
                blocks.append(("paragraph", value))
            continue

        if element.name in {"ul", "ol"}:
            for li in element.find_all("li", recursive=False):
                value = clean_line(li.get_text(" ", strip=True))
                if value:
                    blocks.append(("bullet", value))
            continue

        if element.name == "table":
            rows = element.find_all("tr")
            for row in rows:
                cells = row.find_all(["th", "td"])
                values = [clean_line(cell.get_text(" ", strip=True)) for cell in cells]
                values = [value for value in values if value]
                if values:
                    blocks.append(("paragraph", " | ".join(values)))

    return blocks


def dedupe_preserve_order(items: Iterable[str]) -> list[str]:
    seen: set[str] = set()
    result: list[str] = []
    for item in items:
        if item in seen:
            continue
        seen.add(item)
        result.append(item)
    return result


def dedupe_block_list(blocks: Iterable[tuple[str, str]]) -> list[tuple[str, str]]:
    seen: set[tuple[str, str]] = set()
    result: list[tuple[str, str]] = []
    for block in blocks:
        if block in seen:
            continue
        seen.add(block)
        result.append(block)
    return result


def extract_page_content(html_file: Path) -> list[tuple[str, list[tuple[str, str]]]]:
    soup = BeautifulSoup(html_file.read_text(encoding="utf-8"), "lxml")
    main = soup.select_one("main.content")
    if main is None:
        return []

    boxes = main.find_all(["article", "section"], recursive=False)
    page_sections: list[tuple[str, list[tuple[str, str]]]] = []

    for index, box in enumerate(boxes, start=1):
        heading = box.find(["h1", "h2", "h3"])
        title = clean_line(heading.get_text(" ", strip=True)) if heading else f"Infobox {index}"
        blocks = extract_box_blocks(box, title)
        blocks = dedupe_block_list(blocks)
        if blocks:
            page_sections.append((title, blocks))

    return page_sections


def build_document() -> None:
    html_files = sorted(HTML_DIR.glob("*.html"))

    doc = Document()
    doc.add_heading("TSV Rudow Förderverein - Inhaltsfreigabe", level=0)
    intro = doc.add_paragraph()
    intro.add_run("Zweck: ").bold = True
    intro.add_run(
        "Dieses Dokument enthält die Texte aus den Informationsboxen der Website "
        "zur inhaltlichen Prüfung, Korrektur und Vorschlagsarbeit."
    )
    meta = doc.add_paragraph(
        f"Stand: {datetime.now().strftime('%d.%m.%Y %H:%M')} | Quelle: html/*.html"
    )
    meta.runs[0].font.size = Pt(10)
    web_url = doc.add_paragraph(f"Web URL: {WEBSITE_URL}")
    web_url.runs[0].font.size = Pt(10)

    for html_file in html_files:
        page_title = html_file.stem
        sections = extract_page_content(html_file)
        if not sections:
            continue

        doc.add_page_break()
        doc.add_heading(f"Seite: {page_title}.html", level=1)

        for idx, (section_title, blocks) in enumerate(sections, start=1):
            doc.add_heading(f"Infobox {idx}: {section_title}", level=2)
            for block_type, value in blocks:
                if block_type == "subheading":
                    doc.add_heading(value, level=3)
                elif block_type == "bullet":
                    doc.add_paragraph(value, style="List Bullet")
                else:
                    doc.add_paragraph(value)

    OUTPUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_DOCX)


def build_document_by_topic() -> None:
    html_files = sorted(HTML_DIR.glob("*.html"))
    grouped: dict[str, list[tuple[str, str, list[tuple[str, str]]]]] = {}

    for html_file in html_files:
        stem = html_file.stem
        topic = TOPIC_MAP.get(stem, "Sonstiges")
        sections = extract_page_content(html_file)
        if not sections:
            continue
        grouped.setdefault(topic, [])
        for section_title, blocks in sections:
            grouped[topic].append((stem, section_title, blocks))

    doc = Document()
    doc.add_heading("TSV Rudow Förderverein - Inhaltsfreigabe nach Themen", level=0)
    intro = doc.add_paragraph()
    intro.add_run("Zweck: ").bold = True
    intro.add_run(
        "Dieses Dokument bündelt die Texte aus den Informationsboxen thematisch, "
        "damit die inhaltliche Prüfung und Textvorschläge einfacher abgestimmt werden können."
    )
    meta = doc.add_paragraph(
        f"Stand: {datetime.now().strftime('%d.%m.%Y %H:%M')} | Quelle: html/*.html"
    )
    meta.runs[0].font.size = Pt(10)
    web_url = doc.add_paragraph(f"Web URL: {WEBSITE_URL}")
    web_url.runs[0].font.size = Pt(10)

    for topic in sorted(grouped):
        doc.add_page_break()
        doc.add_heading(f"Thema: {topic}", level=1)
        entries = grouped[topic]

        for page_stem, section_title, blocks in entries:
            doc.add_heading(f"Seite: {page_stem}.html | Infobox: {section_title}", level=2)
            for block_type, value in blocks:
                if block_type == "subheading":
                    doc.add_heading(value, level=3)
                elif block_type == "bullet":
                    doc.add_paragraph(value, style="List Bullet")
                else:
                    doc.add_paragraph(value)

    OUTPUT_DOCX_BY_TOPIC.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_DOCX_BY_TOPIC)


if __name__ == "__main__":
    build_document()
    build_document_by_topic()
